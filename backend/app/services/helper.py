import os
import shutil
from tempfile import NamedTemporaryFile
from fastapi import UploadFile
import fitz
from docx import Document
from langchain.prompts import ChatPromptTemplate, SystemMessagePromptTemplate, HumanMessagePromptTemplate


system_message = SystemMessagePromptTemplate.from_template(
    "You are an intelligent resume parser. Only respond with valid JSON. No markdown. No explanations."
)

human_message = HumanMessagePromptTemplate.from_template(
    """
Extract structured information from the following resume text and return valid, strict JSON format. 
Only include present fields. Do not return nulls, empty strings, or placeholders.
IMPORTANT: Your response must be valid JSON. Use only double quotes. No comments. No trailing commas.

Resume text:
{resume_text}

Expected JSON format:
{{
  "name": "Full name",
  "email": "Email address",
  "phone": "Phone number",
  "linkedin": "LinkedIn profile URL",
  "github": "GitHub profile URL",
  "summary": "Professional summary",
  "skills": ["Skill 1", "Skill 2", "..."],
  "work_experience": [
    {{
      "company": "Company name",
      "position": "Job title",
      "duration": "From – To or Present",
      "location": "City, Country",
      "description": ["Responsibility 1", "Responsibility 2"]
    }}
  ],
  "projects": [
    {{
      "name": "Project name",
      "description": "Short project description",
      "link": "GitHub or live link (optional)"
    }}
  ],
  "education": [
    {{
      "degree": "Degree name",
      "university": "University name",
      "years": "From – To or year of graduation",
      "grade": "GPA/SGPA (optional)"
    }}
  ],
  "certifications": ["Certification 1", "Certification 2"],
  "languages": ["Language 1", "Language 2"],
  "address": "Address if available",
  "interests": ["Interest 1", "Interest 2"]
}}
"""
)

chat_prompt = ChatPromptTemplate.from_messages([system_message, human_message])


def save_upload_to_temp_file(file: UploadFile, suffix=".pdf") -> str:
    """
    Save an uploaded file to a temporary file and return the file path.
    """
    with NamedTemporaryFile(delete=False, suffix=suffix) as temp_file:
        shutil.copyfileobj(file.file, temp_file)
        return temp_file.name


def extract_text(file_path: str) -> str:
    """
    Extracts text from PDF or DOCX resume.
    """
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.pdf':
        return extract_text_adaptive_pdf(file_path)
    elif ext == '.docx':
        return extract_text_docx(file_path)
    else:
        raise ValueError("Unsupported file type. Only .pdf and .docx are allowed.")

def extract_text_adaptive_pdf(file_path: str) -> str:
    """
    Extract text from PDF considering two-column layout if detected.
    """
    doc = fitz.open(file_path)
    if is_two_column_layout(doc):
        left, right = extract_text_by_columns(doc)
        return left + "\n" + right
    else:
        return "\n".join(page.get_text() for page in doc).strip()

def extract_text_docx(file_path: str) -> str:
    """
    Extract text from DOCX file.
    """
    doc = Document(file_path)
    return "\n".join(para.text for para in doc.paragraphs).strip()

def is_two_column_layout(doc) -> bool:
    """
    Heuristic to detect if the PDF has a two-column layout.
    """
    page = doc[0]
    blocks = page.get_text("blocks")
    page_width = page.rect.width
    left_count = right_count = 0

    for block in blocks:
        x0, x1 = block[0], block[2]
        text = block[4].strip()
        if not text:
            continue
        if x1 < page_width / 2:
            left_count += 1
        else:
            right_count += 1

    ratio = min(left_count, right_count) / max(left_count, right_count) if max(left_count, right_count) else 0
    return ratio > 0.5

def extract_text_by_columns(doc) -> tuple:
    """
    Extract text from both left and right columns of a two-column layout.
    """
    left_col_text = ""
    right_col_text = ""
    for page in doc:
        blocks = page.get_text("blocks")
        page_width = page.rect.width
        for block in blocks:
            x0, y0, x1, y1, text, *_ = block
            if not text.strip():
                continue
            if x1 <= page_width / 2:
                left_col_text += text.strip() + "\n"
            else:
                right_col_text += text.strip() + "\n"
    return left_col_text.strip(), right_col_text.strip()
